import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from sample_data import SAMPLE_RESUMES

output_dir = os.path.join(os.path.dirname(__file__), "sample_resumes")
os.makedirs(output_dir, exist_ok=True)

# Generate PDF resumes
styles = getSampleStyleSheet()
normal_style = styles["Normal"]
heading_style = styles["Heading1"]

for name, text in SAMPLE_RESUMES.items():
    clean_name = name.split(" - ")[0].replace(" ", "_").lower()
    
    # 1. Generate PDF
    pdf_path = os.path.join(output_dir, f"{clean_name}_resume.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    
    for line in text.strip().split("\n"):
        if not line.strip():
            story.append(Spacer(1, 8))
        elif line.isupper() and len(line) < 30:
            story.append(Paragraph(f"<b>{line}</b>", styles["Heading2"]))
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(line.replace("&", "&amp;"), normal_style))
            story.append(Spacer(1, 2))
            
    doc.build(story)
    print(f"Created PDF: {pdf_path}")
    
    # 2. Generate DOCX
    docx_path = os.path.join(output_dir, f"{clean_name}_resume.docx")
    doc_docx = docx.Document()
    for line in text.strip().split("\n"):
        if not line.strip():
            continue
        elif line.isupper() and len(line) < 30:
            doc_docx.add_heading(line, level=2)
        else:
            doc_docx.add_paragraph(line)
    doc_docx.save(docx_path)
    print(f"Created DOCX: {docx_path}")

print("All sample resumes generated successfully!")
